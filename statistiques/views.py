"""Vues du module Statistiques : tableau de bord, exports (UC-40, UC-41)."""

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render

from comptes.mixins import roles_requis

from .services import (
    effectifs_par_annee,
    effectifs_par_classe,
    moyennes_par_classe,
    recettes_par_mois,
    tableau_effectifs,
    taux_recouvrement,
)


@login_required
@roles_requis(["ADMIN", "SUPERADMIN", "CENSEUR"])
def dashboard(request):
    """Tableau de bord : effectifs, recouvrement, moyennes (UC-40)."""
    effectifs = effectifs_par_classe()
    evolution = effectifs_par_annee()
    recettes = recettes_par_mois()
    moyennes = moyennes_par_classe()
    recouvrement = taux_recouvrement()
    effectifs_cycles = tableau_effectifs()

    contexte = {
        "effectifs": effectifs,
        "effectifs_json": [{"libelle": e["libelle"], "effectif": e["effectif"]} for e in effectifs],
        "evolution_json": evolution,
        "recettes_json": [{"mois": m, "total": float(t)} for m, t in recettes.items()],
        "moyennes_json": moyennes,
        "recouvrement": recouvrement,
        "effectifs_cycles": effectifs_cycles,
        "total_effectifs": sum(e["effectif"] for e in effectifs),
    }
    return render(request, "statistiques/dashboard.html", contexte)


@login_required
@roles_requis(["ADMIN", "SUPERADMIN", "CENSEUR"])
def export_pdf(request):
    """Rapport PDF des principaux indicateurs."""
    contexte = {
        "effectifs": effectifs_par_classe(),
        "evolution": effectifs_par_annee(),
        "recettes": recettes_par_mois(),
        "moyennes": moyennes_par_classe(),
        "recouvrement": taux_recouvrement(),
        "effectifs_cycles": tableau_effectifs(),
    }
    from documents.pdf import reponse_pdf
    return reponse_pdf("statistiques/rapport_pdf.html", contexte, "rapport_statistiques")


@login_required
@roles_requis(["ADMIN", "SUPERADMIN", "CENSEUR"])
def export_excel(request):
    """Export Excel des effectifs et du recouvrement (UC-41)."""
    import openpyxl
    from openpyxl.styles import Font

    classeur = openpyxl.Workbook()

    # Feuille 1 : effectifs par classe
    feuille = classeur.active
    feuille.title = "Effectifs"
    feuille.append(["Classe", "Effectif"])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for e in effectifs_par_classe():
        feuille.append([e["libelle"], e["effectif"]])

    # Feuille 2 : recouvrement
    feuille2 = classeur.create_sheet("Recouvrement")
    feuille2.append(["Indicateur", "Valeur"])
    for cellule in feuille2[1]:
        cellule.font = Font(bold=True)
    r = taux_recouvrement()
    feuille2.append(["Total dû", float(r["total_du"])])
    feuille2.append(["Total payé", float(r["total_paye"])])
    feuille2.append(["Taux de recouvrement (%)", r["taux"]])

    # Feuille 3 : moyennes par classe
    feuille3 = classeur.create_sheet("Moyennes")
    feuille3.append(["Classe", "Moyenne générale"])
    for cellule in feuille3[1]:
        cellule.font = Font(bold=True)
    for m in moyennes_par_classe():
        feuille3.append([m["libelle"], m["moyenne"]])

    reponse = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    reponse["Content-Disposition"] = 'attachment; filename="rapport_statistiques.xlsx"'
    classeur.save(reponse)
    return reponse
